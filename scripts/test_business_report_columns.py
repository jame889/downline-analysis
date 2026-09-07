import unittest

from openpyxl import Workbook
from import_data import FIRST_GLOBAL_COLUMNS, process_file


class ReportColumnsTest(unittest.TestCase):
    def worksheet(self, order=None):
        order = list(range(27)) if order is None else order
        sheet = Workbook().active
        values = [2, '900057', 'Test member', '2025-01-02', 'SPS', 'SP', None,
                  'TH', '2026-12', 926, 'N', 0, 36, '900001', '900002', 'O',
                  'O', 'SP', 'PS', 6030, 38335, 1575, 32432, 4455, 5903, 0, 0]
        sheet.append([FIRST_GLOBAL_COLUMNS[i][0] if FIRST_GLOBAL_COLUMNS[i][1] != 'ขวา' else None for i in order])
        sheet.append([FIRST_GLOBAL_COLUMNS[i][1] or None for i in order])
        sheet.append([values[i] for i in order])
        return sheet

    def test_september_reorder_preserves_all_fields(self):
        old_members, new_members = {}, {}
        old = process_file(self.worksheet(), '2026-09', old_members)
        # Actual September 7 export order, including moved duplicate position headers.
        order = [0, 1, 2, 5, 15, 16, 12, 17, 18, 19, 20, 21, 22,
                 23, 24, 25, 26, 13, 14, 3, 4, 6, 7, 8, 9, 10, 11]
        new = process_file(self.worksheet(order), '2026-09', new_members)
        self.assertEqual(old, new)
        self.assertEqual(old_members, new_members)
        self.assertEqual(new[0]['current_month_vol_left'], 4455)
        self.assertEqual(new[0]['total_vol_right'], 38335)
        self.assertEqual(new_members['900057']['sponsor_id'], '900001')
        self.assertEqual(new_members['900057']['upline_id'], '900002')

    def test_missing_bv_header_rejected(self):
        sheet = self.worksheet()
        sheet.cell(1, 24).value = 'Unknown BV'
        with self.assertRaisesRegex(ValueError, 'Missing First Global columns'):
            process_file(sheet, '2026-09', {})

    def test_duplicate_side_rejected(self):
        sheet = self.worksheet()
        sheet.cell(2, 25).value = 'ซ้าย'
        with self.assertRaisesRegex(ValueError, 'Duplicate First Global column'):
            process_file(sheet, '2026-09', {})

    def test_missing_relationship_rejected(self):
        sheet = self.worksheet()
        sheet.cell(1, 14).value = 'Unknown relationship'
        with self.assertRaisesRegex(ValueError, 'Missing First Global columns'):
            process_file(sheet, '2026-09', {})


if __name__ == '__main__':
    unittest.main()
